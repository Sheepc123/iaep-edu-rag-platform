"""
测试文档处理功能
"""
import os
import tempfile
from pathlib import Path

def test_document_processor():
    """测试文档处理器"""
    
    print("🧪 测试文档处理功能...")
    
    # 1. 测试依赖库
    print("\n1. 检查依赖库...")
    try:
        import PyPDF2
        print("✅ PyPDF2 已安装")
    except ImportError:
        print("❌ PyPDF2 未安装")
    
    try:
        from docx import Document
        print("✅ python-docx 已安装")
    except ImportError:
        print("❌ python-docx 未安装")
    
    try:
        import fitz
        print("✅ PyMuPDF 已安装")
    except ImportError:
        print("❌ PyMuPDF 未安装")
    
    # 2. 测试临时目录
    print("\n2. 测试临时目录...")
    temp_dir = tempfile.gettempdir()
    print(f"临时目录: {temp_dir}")
    print(f"临时目录存在: {os.path.exists(temp_dir)}")
    print(f"临时目录可写: {os.access(temp_dir, os.W_OK)}")
    
    # 3. 创建测试Word文档
    print("\n3. 创建测试Word文档...")
    try:
        from docx import Document
        
        # 创建一个简单的Word文档
        doc = Document()
        doc.add_heading('测试文档', 0)
        doc.add_paragraph('这是一个测试段落。')
        doc.add_paragraph('包含中文内容的测试。')
        
        # 添加表格
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '姓名'
        table.cell(0, 1).text = '年龄'
        table.cell(1, 0).text = '张三'
        table.cell(1, 1).text = '25'
        
        # 保存到临时文件
        test_docx_path = os.path.join(temp_dir, "test_document.docx")
        doc.save(test_docx_path)
        print(f"✅ 测试Word文档创建成功: {test_docx_path}")
        
        # 测试读取
        test_doc = Document(test_docx_path)
        text_content = ""
        for paragraph in test_doc.paragraphs:
            text_content += paragraph.text + "\n"
        
        print(f"✅ Word文档读取成功，内容长度: {len(text_content)} 字符")
        print(f"内容预览: {text_content[:100]}...")
        
        # 清理
        os.remove(test_docx_path)
        
    except Exception as e:
        print(f"❌ Word文档测试失败: {e}")
    
    # 4. 创建测试PDF文档
    print("\n4. 创建测试PDF文档...")
    try:
        # 尝试使用reportlab创建PDF
        try:
            from reportlab.pdfgen import canvas
            from reportlab.lib.pagesizes import letter
            
            test_pdf_path = os.path.join(temp_dir, "test_document.pdf")
            c = canvas.Canvas(test_pdf_path, pagesize=letter)
            c.drawString(100, 750, "测试PDF文档")
            c.drawString(100, 720, "这是一个测试PDF文档。")
            c.drawString(100, 690, "包含中文内容的测试。")
            c.save()
            
            print(f"✅ 测试PDF文档创建成功: {test_pdf_path}")
            
            # 测试读取
            import fitz
            doc = fitz.open(test_pdf_path)
            text_content = ""
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text_content += page.get_text()
            doc.close()
            
            print(f"✅ PDF文档读取成功，内容长度: {len(text_content)} 字符")
            print(f"内容预览: {text_content[:100]}...")
            
            # 清理
            os.remove(test_pdf_path)
            
        except ImportError:
            print("⚠️  reportlab未安装，跳过PDF创建测试")
            
    except Exception as e:
        print(f"❌ PDF文档测试失败: {e}")
    
    # 5. 测试文件路径处理
    print("\n5. 测试文件路径处理...")
    
    # 测试中文文件名
    chinese_filename = "测试文档_中文名称.docx"
    safe_filename = f"temp_{hash(chinese_filename) % 1000000}.docx"
    print(f"原始文件名: {chinese_filename}")
    print(f"安全文件名: {safe_filename}")
    
    # 测试UUID文件名
    import uuid
    uuid_filename = f"temp_{uuid.uuid4().hex}.docx"
    print(f"UUID文件名: {uuid_filename}")
    
    print("\n🎉 文档处理测试完成！")
    
    print("\n💡 建议:")
    print("1. 确保所有依赖库都已安装")
    print("2. 使用UUID生成临时文件名避免中文路径问题")
    print("3. 添加详细的错误日志帮助调试")
    print("4. 检查文件权限和临时目录访问权限")

if __name__ == "__main__":
    test_document_processor()
