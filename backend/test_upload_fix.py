"""
测试文件上传修复
"""
import requests
import os
from datetime import datetime

def test_upload_fix():
    """测试上传修复"""
    
    print("🧪 测试文件上传修复...")
    
    # 1. 登录获取token
    print("\n1. 用户登录...")
    login_data = {
        "username": "teacher1",
        "password": "123456"
    }
    
    try:
        response = requests.post("http://127.0.0.1:8000/api/v1/auth/login", json=login_data)
        if response.status_code == 200:
            token_data = response.json()
            access_token = token_data["access_token"]
            print("✅ 登录成功")
            
            headers = {
                "Authorization": f"Bearer {access_token}"
            }
        else:
            print(f"❌ 登录失败: {response.status_code}")
            return
    except Exception as e:
        print(f"❌ 登录异常: {e}")
        return
    
    # 2. 创建测试Word文档
    print("\n2. 创建测试Word文档...")
    try:
        from docx import Document
        
        # 创建一个简单的Word文档
        doc = Document()
        doc.add_heading('测试文档标题', 0)
        doc.add_paragraph('这是第一个测试段落，包含一些中文内容。')
        doc.add_paragraph('这是第二个测试段落，用于验证文档处理功能。')
        
        # 添加表格
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = '项目'
        table.cell(0, 1).text = '内容'
        table.cell(1, 0).text = '测试类型'
        table.cell(1, 1).text = '文档上传测试'
        table.cell(2, 0).text = '时间'
        table.cell(2, 1).text = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        
        # 保存测试文档
        test_file_path = "test_upload_document.docx"
        doc.save(test_file_path)
        print(f"✅ 测试Word文档创建成功: {test_file_path}")
        
        # 检查文件大小
        file_size = os.path.getsize(test_file_path)
        print(f"文件大小: {file_size} bytes")
        
        if file_size == 0:
            print("❌ 创建的文档文件为空")
            return
        
    except Exception as e:
        print(f"❌ 创建测试文档失败: {e}")
        return
    
    # 3. 测试文档上传
    print("\n3. 测试文档上传...")
    try:
        with open(test_file_path, 'rb') as f:
            files = {
                'file': ('test_upload_document.docx', f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            }
            data = {
                'category': '测试文档',
                'tags': '测试,上传,修复'
            }
            
            response = requests.post(
                "http://127.0.0.1:8000/api/v1/teacher-knowledge/upload",
                headers=headers,
                files=files,
                data=data
            )
            
            print(f"上传响应状态: {response.status_code}")
            
            if response.status_code == 200:
                result = response.json()
                print(f"✅ 文档上传成功!")
                print(f"文档ID: {result.get('document_id')}")
                print(f"消息: {result.get('message')}")
            else:
                print(f"❌ 文档上传失败: {response.status_code}")
                print(f"错误信息: {response.text}")
                
    except Exception as e:
        print(f"❌ 上传测试异常: {e}")
    finally:
        # 清理测试文件
        if os.path.exists(test_file_path):
            os.remove(test_file_path)
            print(f"🧹 清理测试文件: {test_file_path}")
    
    # 4. 检查上传目录
    print("\n4. 检查上传目录...")
    upload_base_dir = "uploads/knowledge_base"
    current_date = datetime.now()
    current_month_dir = os.path.join(
        upload_base_dir,
        f"{current_date.year:04d}",
        f"{current_date.month:02d}"
    )
    
    print(f"预期上传目录: {current_month_dir}")
    
    if os.path.exists(current_month_dir):
        files = os.listdir(current_month_dir)
        print(f"✅ 上传目录存在，包含 {len(files)} 个文件")
        if files:
            print("文件列表:")
            for file in files[:5]:  # 只显示前5个文件
                file_path = os.path.join(current_month_dir, file)
                file_size = os.path.getsize(file_path)
                print(f"  - {file} ({file_size} bytes)")
    else:
        print(f"⚠️  上传目录不存在: {current_month_dir}")
    
    print("\n🎉 测试完成!")

if __name__ == "__main__":
    test_upload_fix()
